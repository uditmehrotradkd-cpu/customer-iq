"""Build the Team Doomsday presentation (PowerPoint) from the project's results and figures.

Usage:  python make_deck.py [output.pptx]      (default: ~/Downloads/Doomsday_CustomerIQ.pptx)
Needs:  pip install python-pptx pillow
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "artifacts" / "figures"
MEDIA = ROOT / "web" / "static" / "media"

# Doomsday palette: void black, blood red, ember orange, ash text.
BG, PANEL, PANEL_2, EDGE = "0A0A0C", "16161A", "1E1E23", "2C2C32"
RED, EMBER, WHITE, ASH, DIM = "C1121F", "FF5A1F", "F5F1EA", "BDB7AE", "77716B"
TITLE_FONT, BODY_FONT = "Bahnschrift", "Segoe UI"
TOTAL = 23
W, H = Inches(13.333), Inches(7.5)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_alpha(shape, percent: float) -> None:
    srgb = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(percent * 1000))}))


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE, rotation=0.0):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.rotation = rotation
    return s


def _runs(paragraph, text: str, size: float, color: str, bold: bool, font: str, spacing: int | None, italic: bool = False):
    """Add runs to a paragraph; **text** is rendered bold and white."""
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        strong = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run()
        run.text = part[2:-2] if strong else part
        run.font.size = Pt(size)
        run.font.name = font
        run.font.bold = bold or strong
        run.font.italic = italic
        run.font.color.rgb = rgb(WHITE if strong and color == ASH else color)
        if spacing:
            run._r.get_or_add_rPr().set("spc", str(spacing))


def text(slide, x, y, w, h, content, size=16, color=ASH, bold=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, italic=False, line_spacing=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(content.split("\n") if isinstance(content, str) else content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        _runs(p, line, size, color, bold, font, spacing, italic)
    return box


def bullets(slide, x, y, w, h, items, size=16, gap=8):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        sub = isinstance(item, tuple)
        line = item[0] if sub else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap if not sub else gap - 3)
        marker = p.add_run()
        marker.text = ("      –  " if sub else "■  ")
        marker.font.size = Pt(size * (0.9 if sub else 0.6))
        marker.font.color.rgb = rgb(EMBER if sub else RED)
        marker.font.name = BODY_FONT
        _runs(p, line, size - (2 if sub else 0), ASH, False, BODY_FONT, None)
    return box


def picture(slide, path: Path, x, y, w, h, frame=True):
    """Fit an image inside the box (keeps aspect ratio), optionally with a blood-red frame."""
    if not path.exists():
        return None
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    px, py = x + (w - pw) // 2, y + (h - ph) // 2
    if frame:
        pad = Inches(0.06)
        rect(slide, px - pad, py - pad, pw + 2 * pad, ph + 2 * pad, fill=PANEL, line=RED, line_w=1.25)
    return slide.shapes.add_picture(str(path), px, py, pw, ph)


def cover_picture(slide, path: Path, x, y, w, h):
    """Fill the box with the image, cropping the overflow (like CSS object-fit: cover)."""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        crop = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = crop
    else:
        crop = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = crop
    return pic


def table(slide, x, y, w, rows, col_widths=None, size=12, row_h=0.38, highlight_first_col=True, bold_rows=()):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, Inches(row_h * n_rows))
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set("firstRow", "0")
    tbl_pr.set("bandRow", "0")
    t = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            t.columns[i].width = int(w * cw / total)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(RED if r == 0 else (PANEL if r % 2 else PANEL_2))
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tc_pr = cell._tc.get_or_add_tcPr()
            # Border lines must precede the cell fill in tcPr.
            for pos, edge in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                ln = tc_pr.makeelement(qn(edge), {"w": "9525"})
                fill = ln.makeelement(qn("a:solidFill"), {})
                fill.append(fill.makeelement(qn("a:srgbClr"), {"val": EDGE}))
                ln.append(fill)
                tc_pr.insert(pos, ln)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            header = r == 0
            strong = header or r in bold_rows or (highlight_first_col and c == 0)
            _runs(p, str(value), size - (1 if header else 0), WHITE if strong else ASH, strong, TITLE_FONT if header else BODY_FONT,
                  80 if header else None)
    return shape


def chrome(slide, idx: int, title: str, kicker: str = "CUSTOMERIQ"):
    """Shared Doomsday frame: red edge, kicker, countdown, title rule and footer."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(BG)
    rect(slide, 0, 0, Inches(0.12), H, fill=RED)
    text(slide, Inches(0.55), Inches(0.32), Inches(8), Inches(0.3), f"TEAM DOOMSDAY  //  {kicker}", size=10, color=RED, bold=True,
         font=TITLE_FONT, spacing=300)
    left = TOTAL - idx
    text(slide, Inches(10.4), Inches(0.26), Inches(2.5), Inches(0.45), f"T–{left:02d}", size=22, color=EMBER, bold=True, font=TITLE_FONT,
         align=PP_ALIGN.RIGHT, spacing=200)
    text(slide, Inches(0.55), Inches(0.62), Inches(11.8), Inches(0.7), title.upper(), size=30, color=WHITE, bold=True, font=TITLE_FONT,
         spacing=50)
    rect(slide, Inches(0.55), Inches(1.33), Inches(1.1), Inches(0.06), fill=RED)
    rect(slide, Inches(1.7), Inches(1.33), Inches(0.35), Inches(0.06), fill=EMBER)
    rect(slide, Inches(0.55), Inches(7.02), Inches(12.35), Emu(9525), fill=EDGE)
    text(slide, Inches(0.55), Inches(7.1), Inches(9), Inches(0.25), "CUSTOMER SEGMENTATION  ·  FROM RAW DATA TO RESPONSIBLE ACTION", size=9,
         color=DIM, font=TITLE_FONT, spacing=150)
    text(slide, Inches(10.9), Inches(7.1), Inches(2), Inches(0.25), f"{idx:02d} / {TOTAL:02d}", size=9, color=DIM, font=TITLE_FONT,
         align=PP_ALIGN.RIGHT, spacing=150)


def notes(slide, content: str) -> None:
    slide.notes_slide.notes_text_frame.text = content


def card(slide, x, y, w, h, heading, body, accent=RED, size=14):
    rect(slide, x, y, w, h, fill=PANEL, line=EDGE)
    rect(slide, x, y, w, Inches(0.07), fill=accent)
    text(slide, x + Inches(0.25), y + Inches(0.25), w - Inches(0.5), Inches(0.4), heading.upper(), size=13, color=accent, bold=True,
         font=TITLE_FONT, spacing=200)
    if isinstance(body, list):
        bullets(slide, x + Inches(0.25), y + Inches(0.75), w - Inches(0.5), h - Inches(0.9), body, size=size, gap=6)
    else:
        text(slide, x + Inches(0.25), y + Inches(0.75), w - Inches(0.5), h - Inches(0.9), body, size=size)


def stat(slide, x, y, w, value, label, color=EMBER):
    rect(slide, x, y, w, Inches(1.2), fill=PANEL, line=EDGE)
    rect(slide, x, y, Inches(0.06), Inches(1.2), fill=color)
    text(slide, x + Inches(0.25), y + Inches(0.12), w - Inches(0.3), Inches(0.6), value, size=30, color=WHITE, bold=True, font=TITLE_FONT)
    text(slide, x + Inches(0.25), y + Inches(0.78), w - Inches(0.3), Inches(0.3), label.upper(), size=10, color=color, bold=True,
         font=TITLE_FONT, spacing=200)


def arrow(slide, x1, y1, x2, y2, color=EMBER, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.75)
    if head:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def doomsday_clock(slide, cx, cy, r):
    """Clock face at two minutes to midnight, drawn with native shapes."""
    import math

    for scale, alpha in ((1.55, 8), (1.3, 14)):
        glow = rect(slide, int(cx - r * scale), int(cy - r * scale), int(2 * r * scale), int(2 * r * scale), fill=RED, shape=MSO_SHAPE.OVAL)
        set_alpha(glow, alpha)
    rect(slide, int(cx - r), int(cy - r), int(2 * r), int(2 * r), fill="111114", line=EMBER, line_w=3, shape=MSO_SHAPE.OVAL)
    for hour in range(12):
        a = math.radians(hour * 30)
        long_tick = hour % 3 == 0
        tw, th = Inches(0.07 if long_tick else 0.04), Inches(0.34 if long_tick else 0.2)
        dist = r - Inches(0.3)
        tx, ty = cx + dist * math.sin(a), cy - dist * math.cos(a)
        rect(slide, int(tx - tw / 2), int(ty - th / 2), tw, th, fill=RED if hour == 0 else ASH, rotation=hour * 30)

    def hand(angle_deg, length, width, color):
        a = math.radians(angle_deg)
        mx, my = cx + (length / 2) * math.sin(a), cy - (length / 2) * math.cos(a)
        rect(slide, int(mx - width / 2), int(my - length / 2), int(width), int(length), fill=color, rotation=angle_deg)

    hand((11 + 58 / 60) * 30, r * 0.5, Inches(0.11), WHITE)
    hand(58 * 6, r * 0.78, Inches(0.07), WHITE)
    hand(0, r * 0.82, Inches(0.025), RED)
    dot = Inches(0.22)
    rect(slide, int(cx - dot / 2), int(cy - dot / 2), dot, dot, fill=RED, shape=MSO_SHAPE.OVAL)


# ---------------------------------------------------------------------- slides
def build(out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    new = lambda: prs.slides.add_slide(blank)  # noqa: E731

    # 1. Title
    s = new()
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0, 0, Inches(0.12), H, fill=RED)
    rect(s, 0, Inches(7.38), W, Inches(0.12), fill=RED)
    text(s, Inches(0.8), Inches(0.75), Inches(7), Inches(0.4), "TEAM DOOMSDAY PRESENTS", size=14, color=RED, bold=True, font=TITLE_FONT, spacing=500)
    text(s, Inches(0.8), Inches(1.3), Inches(7.5), Inches(2.4), "CUSTOMER\nSEGMENTATION", size=60, color=WHITE, bold=True, font=TITLE_FONT,
         spacing=100, line_spacing=0.9)
    rect(s, Inches(0.82), Inches(3.55), Inches(1.4), Inches(0.08), fill=RED)
    rect(s, Inches(2.27), Inches(3.55), Inches(0.45), Inches(0.08), fill=EMBER)
    text(s, Inches(0.8), Inches(3.85), Inches(7), Inches(0.9), "From raw customer data to responsible business actions", size=22, color=ASH)
    text(s, Inches(0.8), Inches(4.75), Inches(7), Inches(0.5), "11:58 — two minutes to midnight for one-size-fits-all marketing.", size=15,
         color=EMBER, italic=True)
    text(s, Inches(0.8), Inches(5.9), Inches(7.5), Inches(0.8),
         ["Segmentation model  ·  Jupyter notebook  ·  Streamlit dashboard", "Production website + REST API  ·  AI data agent  ·  reusable Python package"],
         size=12, color=DIM, font=TITLE_FONT, spacing=100)
    doomsday_clock(s, Inches(10.2), Inches(3.6), Inches(2.1))
    text(s, Inches(8.6), Inches(6.35), Inches(3.2), Inches(0.4), "T–22  ·  COUNTDOWN STARTED", size=11, color=RED, bold=True, font=TITLE_FONT,
         align=PP_ALIGN.CENTER, spacing=300)
    notes(s, "Every customer is different. Some buy often, some spend a lot, some only show up for special offers. "
             "Team Doomsday found those groups in the data and explains how to serve each one responsibly.")

    # 2. Problem
    s = new()
    chrome(s, 2, "Business problem and objective", "THE THREAT")
    card(s, Inches(0.55), Inches(1.75), Inches(3.9), Inches(4.9), "Problem",
         "One-size-fits-all marketing wastes budget and irritates customers. The same discount goes to people who would have paid full price "
         "and to people who will never buy.", size=16)
    card(s, Inches(4.7), Inches(1.75), Inches(3.9), Inches(4.9), "Objective",
         "Group customers into meaningful segments and explain how each one should be served: a clear goal, next best actions, KPIs and "
         "guardrails per segment.", accent=EMBER, size=16)
    card(s, Inches(8.85), Inches(1.75), Inches(4.05), Inches(4.9), "Questions answered", [
        "How many distinct customer groups exist, and **why that number**?",
        "What defines each group: spend, frequency, recency, channel, engagement?",
        "Which actions should each group get, and **what should we avoid**?",
    ], size=15)
    notes(s, "The goal is decisions, not just clusters. Every segment comes with a clear goal, next best actions and KPIs.")

    # 3. Dataset
    s = new()
    chrome(s, 3, "Dataset overview", "THE EVIDENCE")
    bullets(s, Inches(0.55), Inches(1.75), Inches(5.9), Inches(5), [
        "**2,240 customers**, 2 years of purchase history (wine, meat, fish, fruit, sweets, gold)",
        "**Spending:** amount spent on 6 product categories",
        "**Frequency & channel:** web, catalog, store and discounted purchases; monthly web visits",
        "**Recency:** days since the last purchase",
        "**Engagement:** acceptance of 6 marketing campaigns, complaints",
        "**Demographics:** birth year, income, children, education, marital status, enrolment date",
        "Arrived **one-hot encoded: 699 columns**, 662 of them dummies for the enrolment date alone",
    ], size=15)
    picture(s, FIG / "01_behaviour_overview.png", Inches(6.75), Inches(1.75), Inches(6.1), Inches(5.0))
    notes(s, "Clustering on hundreds of sparse date dummies would be meaningless, so the loader first collapses the dummy columns back "
             "into categorical values and a proper date.")

    # 4. Data quality issues
    s = new()
    chrome(s, 4, "Data quality issues found", "THE DAMAGE")
    table(s, Inches(0.55), Inches(1.75), Inches(8.1), [
        ["Issue", "Records affected"],
        ["Same profile under a different ID (re-registrations)", "183"],
        ["Birth year before 1920 (implausible age)", "3"],
        ["Income of 666,666 (data-entry error)", "1"],
        ["Spend recorded with zero purchases", "6"],
        ["Deal purchases > total purchases", "1 (capped)"],
        ["Messy categories (Alone, Absurd, YOLO)", "7"],
        ["Constant columns (Z_CostContact, Z_Revenue)", "2 columns"],
    ], col_widths=[3, 1.2], size=14, row_h=0.52, highlight_first_col=False)
    stat(s, Inches(9.1), Inches(1.75), Inches(3.8), "2,240", "raw rows", RED)
    stat(s, Inches(9.1), Inches(3.15), Inches(3.8), "2,047", "clean customers", EMBER)
    stat(s, Inches(9.1), Inches(4.55), Inches(3.8), "8.6%", "removed, all logged", ASH)
    text(s, Inches(0.55), Inches(6.2), Inches(8.1), Inches(0.6), "Every step is written to an **audit trail**, so the cleaning can be checked and repeated.",
         size=14)
    notes(s, "Duplicates would double-count some behaviours and bias the centroids. Every rule is written down and logged.")

    # 5. Cleaning policy
    s = new()
    chrome(s, 5, "Cleaning policy", "DAMAGE CONTROL")
    bullets(s, Inches(0.55), Inches(1.7), Inches(8.4), Inches(5.3), [
        "**Duplicates:** exact copies, re-used IDs and re-registered profiles removed; missing IDs re-assigned",
        "**Blank rows:** more than half the fields missing → dropped, never imputed into fake customers",
        "**Text numbers:** $58,138 · 58,138.00 · 58138 USD converted; N/A, unknown, negatives cleared",
        "**Yes/No flags:** Yes · TRUE · y → 1/0; out-of-range values (2, 5, −1) cleared",
        "**Dates:** 21-08-2013 · 2013-08-21 · 21 Aug 2013 · August 21, 2013 all understood; impossible and future dates cleared",
        "**Impossible records:** born before 1920 or younger than 16, income > 200k, spend without purchases",
        "**Categories:** ' phd.' → PhD, Masters/MBA → Master, Alone → Single, YOLO → Unknown",
        "**Missing values:** income by education median, everything else by median",
        "**Scoring safety:** the live model applies the same rules to every new customer",
    ], size=13.5, gap=7)
    steps = ["RAW FILE", "STANDARDISE", "DROP BLANKS", "DE-DUPLICATE", "HARMONISE", "VALIDATE", "IMPUTE", "CLEAN DATA"]
    for i, step in enumerate(steps):
        y = Inches(1.7 + i * 0.66)
        color = RED if i in (0, len(steps) - 1) else PANEL
        rect(s, Inches(9.6), y, Inches(3.3), Inches(0.46), fill=color, line=RED if color == PANEL else None)
        text(s, Inches(9.6), y, Inches(3.3), Inches(0.46), step, size=12, color=WHITE, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, spacing=200)
        if i < len(steps) - 1:
            arrow(s, Inches(11.25), y + Inches(0.46), Inches(11.25), y + Inches(0.66))
    notes(s, "Real exports are messy, so every rule is built into the pipeline and not done as a one-off step. "
             "The data agent tells the user exactly what it cleaned before it trains.")

    # 6. Stress test
    s = new()
    chrome(s, 6, "Stress test: 1 lakh uncleaned customers", "SURVIVAL TEST")
    table(s, Inches(0.55), Inches(1.7), Inches(7.6), [
        ["Planted problem", "Handled"],
        ["Blank rows", "200 removed"],
        ["Exact duplicate rows", "3,000 removed"],
        ["Re-used customer IDs", "1,000 removed"],
        ["Re-registered profiles under a new ID", "995 removed"],
        ["Impossible birth years (1893, 2031 …)", "378 removed"],
        ["Data-entry incomes (9,999,999 …)", "188 removed"],
        ["Spend recorded with zero purchases", "284 removed"],
        ["Numbers and yes/no flags stored as text", "4,901 converted"],
        ["Inconsistent category spellings", "9,189 fixed"],
    ], col_widths=[2.6, 1.3], size=13, row_h=0.47, highlight_first_col=False)
    stat(s, Inches(8.6), Inches(1.7), Inches(2.05), "1,00,000", "rows in", RED)
    stat(s, Inches(10.85), Inches(1.7), Inches(2.05), "93,955", "clean out", EMBER)
    stat(s, Inches(8.6), Inches(3.1), Inches(2.05), "~40 s", "clean + train", ASH)
    stat(s, Inches(10.85), Inches(3.1), Inches(2.05), "4", "same segments", ASH)
    text(s, Inches(8.6), Inches(4.6), Inches(4.3), Inches(2.2), [
        "A deliberately messy Excel export with every problem planted on purpose, uploaded to the data agent.",
        "",
        "Before the test the cleaner lost **16,000 dates**, **3,000 incomes** and **4,700 categories**. Each gap became a new rule with an automated test.",
    ], size=13)
    notes(s, "The generator is tests/messy_data.py. Before this test the cleaner silently lost 16,000 dates in other formats, 3,000 incomes "
             "written as text and 4,700 misspelled categories, and turned blank rows into fake customers.")

    # 7. Features
    s = new()
    chrome(s, 7, "Feature engineering: measuring behaviour", "THE SIGNALS")
    table(s, Inches(0.55), Inches(1.7), Inches(12.35), [
        ["Feature", "Meaning", "Lens"],
        ["Total_Spend", "Spend across 6 categories", "Monetary"],
        ["Total_Purchases", "Web + catalog + store purchases", "Frequency"],
        ["Recency", "Days since last purchase", "Recency"],
        ["Avg_Order_Value", "Spend ÷ purchases", "Basket size"],
        ["Deal_Ratio", "Discounted ÷ total purchases", "Promotion sensitivity"],
        ["Web_Share, Catalog_Share", "Channel mix", "Channel preference"],
        ["NumWebVisitsMonth", "Monthly web visits", "Digital engagement"],
        ["Campaigns_Accepted", "Campaigns accepted (0–6)", "Marketing responsiveness"],
        ["Income", "Household income", "Affordability"],
    ], col_widths=[1.4, 2, 1.5], size=13, row_h=0.43)
    text(s, Inches(0.55), Inches(6.3), Inches(12.3), Inches(0.5),
         "Demographics (age, children, education, marital status) are used **only to describe** segments, **never to form them**.", size=15)
    notes(s, "Clustering on behaviour rather than identity gives segments we can act on and lowers the risk of discriminatory targeting.")

    # 8. EDA
    s = new()
    chrome(s, 8, "Key insights from exploratory analysis", "THE PATTERNS")
    bullets(s, Inches(0.55), Inches(1.75), Inches(5.9), Inches(5), [
        "**Wine (50%) and meat (28%)** account for most of the spend",
        "**Spend is concentrated:** the top 20% of customers generate **52%** of revenue",
        "**Income drives spend** (ρ = 0.85); **children reduce spend** (ρ = −0.50)",
        "**More web visits ≠ more value:** visits correlate negatively with spend (ρ = −0.47)",
        "**72% never accepted a campaign:** responsiveness sits in a small group",
        "**Recency is uniform (0–99 days):** a lifecycle trigger, not a segment driver",
    ], size=15)
    picture(s, FIG / "04_correlation_heatmap.png", Inches(6.75), Inches(1.7), Inches(6.1), Inches(5.1))
    notes(s, "Heavy browsing combined with low spend points to friction or price shopping, not loyalty.")

    # 9. Preprocessing
    s = new()
    chrome(s, 9, "Preprocessing: why each step matters", "THE ARMOUR")
    bullets(s, Inches(0.55), Inches(1.75), Inches(5.9), Inches(5), [
        "**1. Median imputation:** any future customer can be scored",
        "**2. Winsorising (1st–99th percentile):** extreme spenders cannot drag the centroids",
        "**3. log1p on skewed features** (|skew| > 0.75)",
        ("Avg order value skew drops from **20.8 → −0.1**",),
        "**4. Standard scaling:** income and ratios end up on the same scale",
        "All steps sit in **one sklearn pipeline**: training and scoring apply identical transformations",
    ], size=15)
    picture(s, FIG / "03_skew_correction.png", Inches(6.75), Inches(1.7), Inches(6.1), Inches(5.1))
    notes(s, "K-Means uses Euclidean distance. Without scaling it would cluster almost entirely on income.")

    # 10. Choosing k
    s = new()
    chrome(s, 10, "Choosing the number of segments", "THE VERDICT")
    bullets(s, Inches(0.55), Inches(1.75), Inches(5.9), Inches(5), [
        "Evaluated **k = 2 … 10** on elbow, silhouette, Calinski-Harabasz, Davies-Bouldin and bootstrap stability",
        "**Rule:** elbow k if actionable (k ≥ 3), no segment < 5%, stability ARI ≥ 0.80, silhouette ≥ 90% of best",
        "**Result: k = 4**",
        ("Elbow of the inertia curve at k = 4",),
        ("Silhouette 0.216 (k = 3: 0.234)",),
        ("**Stability ARI 0.97** on 80% subsamples",),
        ("Smallest segment: 16% of customers",),
        "k = 2 only splits high vs. low value: too coarse to act on",
    ], size=15, gap=6)
    picture(s, FIG / "05_k_selection.png", Inches(6.75), Inches(1.7), Inches(6.1), Inches(5.1))
    notes(s, "k = 4 separates premium customers who respond to campaigns from affluent customers who ignore them. "
             "Those two groups need opposite strategies, and k = 3 merges them.")

    # 11. Algorithms
    s = new()
    chrome(s, 11, "Algorithm comparison", "THE CONTENDERS")
    table(s, Inches(0.55), Inches(1.75), Inches(7.3), [
        ["Algorithm", "Silhouette", "Davies-Bouldin", "Smallest", "Scores new"],
        ["K-Means (selected)", "0.216", "1.62", "16.1%", "Yes"],
        ["Agglomerative (Ward)", "0.175", "1.86", "18.2%", "No"],
        ["Gaussian Mixture", "0.023", "1.86", "5.1%", "Yes"],
    ], col_widths=[2, 1, 1.2, 1, 1], size=13, row_h=0.55, bold_rows=(1,))
    bullets(s, Inches(0.55), Inches(4.2), Inches(7.3), Inches(2.6), [
        "**K-Means wins on separation** and is easy to explain: centroid = typical customer",
        "It **assigns new customers in real time**; Ward cannot without a retrain",
        "Gaussian Mixture produced overlapping, unbalanced groups",
    ], size=15)
    picture(s, FIG / "06_silhouette_diagram.png", Inches(8.2), Inches(1.7), Inches(4.7), Inches(5.1))
    notes(s, "Ward clustering is competitive but cannot score a new customer without a retrain.")

    # 12. Four segments
    s = new()
    chrome(s, 12, "Four customer segments at a glance", "THE FACTIONS")
    table(s, Inches(0.55), Inches(1.75), Inches(12.35), [
        ["Segment", "Customers", "Revenue", "Median income", "Avg spend", "Deal ratio", "Campaigns"],
        ["Premium Campaign Responders", "16.1%", "38.7%", "$77.4k", "$1,464", "8%", "1.93"],
        ["Affluent Quiet Loyalists", "24.3%", "39.4%", "$68.4k", "$987", "11%", "0.01"],
        ["Digital Deal-Seekers", "24.1%", "18.5%", "$49.5k", "$467", "30%", "0.38"],
        ["Budget-Conscious Browsers", "35.6%", "3.5%", "$32.0k", "$59", "37%", "0.14"],
    ], col_widths=[2.6, 1, 1, 1.2, 1, 1, 1], size=14, row_h=0.6)
    stat(s, Inches(0.55), Inches(5.1), Inches(4), "40% → 78%", "of customers  →  of revenue", RED)
    text(s, Inches(4.9), Inches(5.15), Inches(8), Inches(1.4), [
        "The two premium segments (40% of customers) generate **78% of revenue**.",
        "The largest segment brings in the least: this is where prioritisation starts.",
    ], size=16)
    notes(s, "The largest segment brings in the least revenue and the smallest segment brings in almost 40%.")

    # 13. Fingerprints
    s = new()
    chrome(s, 13, "Segment fingerprints", "THE PROFILES")
    bullets(s, Inches(0.55), Inches(1.75), Inches(5.9), Inches(5), [
        "Each cell: standard deviations above (red) or below (blue) the average customer",
        "**Premium Campaign Responders:** campaigns +1.7 sd, spend +1.4 sd, income +1.1 sd",
        "**Affluent Quiet Loyalists:** frequent catalog buyers, rare web visitors, ignore campaigns",
        "**Digital Deal-Seekers:** highest web share (44%), frequent visits, deal-driven",
        "**Budget-Conscious Browsers:** fewest purchases (−1.1 sd), highest deal ratio, in-store",
    ], size=15)
    picture(s, FIG / "09_segment_fingerprint.png", Inches(6.75), Inches(1.7), Inches(6.1), Inches(5.1))
    notes(s, "Segment names are not hand-picked: each fingerprint is matched to a library of business archetypes with the Hungarian algorithm.")

    # 14. Channels
    s = new()
    chrome(s, 14, "Channel and product preferences", "THE TERRITORY")
    table(s, Inches(0.55), Inches(1.75), Inches(6.6), [
        ["Segment", "Web", "Catalog", "Store", "Over-indexed"],
        ["Premium Campaign Responders", "29%", "30%", "41%", "Wine, Meat"],
        ["Affluent Quiet Loyalists", "27%", "26%", "48%", "Fruit, Sweets"],
        ["Digital Deal-Seekers", "44%", "13%", "43%", "Wine, Gold"],
        ["Budget-Conscious Browsers", "32%", "6%", "62%", "Gold, Fish"],
    ], col_widths=[2.4, 0.8, 0.9, 0.8, 1.4], size=13, row_h=0.6)
    text(s, Inches(0.55), Inches(5.2), Inches(6.6), Inches(1.5),
         "Channel mix tells us **where to reach** each segment. Category over-indexing tells us **which products to feature** in personalised offers.",
         size=15)
    picture(s, FIG / "12_channel_category_mix.png", Inches(7.45), Inches(1.7), Inches(5.45), Inches(5.1))
    notes(s, "Channel mix tells us where to reach each segment; category over-indexing tells us which products to feature.")

    # 15-18. Segment playbooks
    playbooks = [
        ("Premium Campaign Responders", "16% customers · 39% revenue", "persona-premium.jpg",
         "High income, high spend, large baskets, respond to campaigns", "Retain and grow share of wallet",
         ["Tiered VIP / loyalty programme with early access", "Priority audience for new campaign pilots",
          "Personalised wine and meat offers instead of discounts", "Concierge-style service on catalog and web"],
         "Retention · campaign conversion · share of wallet", "Cap contact frequency: responsiveness is not consent to unlimited messaging"),
        ("Affluent Quiet Loyalists", "24% customers · 39% revenue", "persona-loyalist.jpg",
         "Affluent, frequent buyers who ignore promotional campaigns", "Protect value without over-marketing",
         ["Cut generic campaigns; service- and quality-led messages", "Curated bundles and new-arrival notes",
          "Non-monetary recognition: priority delivery, events", "Watch recency: silent high-value customers churn quietly"],
         "Revenue retention · average order value · unsubscribe rate", "Treat low engagement as a preference; no pressure tactics"),
        ("Digital Deal-Seekers", "24% customers · 19% revenue", "persona-digital.jpg",
         "Mid-income, web-first, browse often, buy on promotion", "Convert browsing into full-price, larger baskets (main growth opportunity)",
         ["Personalised web/app bundles that raise basket size", "Loyalty points instead of blanket discounts",
          "Cart-abandonment and price-drop alerts with honest expiry", "Free-shipping threshold just above current order value"],
         "Web conversion · average order value · full-price share", "No fake scarcity or countdown timers; show genuine savings"),
        ("Budget-Conscious Browsers", "36% customers · 3.5% revenue", "persona-budget.jpg",
         "Lower income, lowest spend, frequent visits, mostly in-store", "Everyday value and habit at low cost to serve",
         ["Value packs and essentials, not premium ranges", "In-store and low-cost email/app; no printed catalogs",
          "Entry-level loyalty tier rewarding visit → purchase", "Simplify the web journey: many visits, few purchases = friction"],
         "Visit-to-purchase conversion · frequency · cost to serve", "Never push credit or upsell beyond what they can afford"),
    ]
    for i, (name, share, photo, who, goal, actions, kpis, guardrail) in enumerate(playbooks):
        idx = 15 + i
        s = new()
        chrome(s, idx, f"Segment {i + 1}: {name}", "FACTION PLAYBOOK")
        cover_picture(s, MEDIA / photo, Inches(0.55), Inches(1.7), Inches(4.3), Inches(5.1))
        shade = rect(s, Inches(0.55), Inches(5.6), Inches(4.3), Inches(1.2), fill=BG)
        set_alpha(shade, 75)
        rect(s, Inches(0.55), Inches(1.7), Inches(4.3), Inches(5.1), line=RED, line_w=1.25)
        text(s, Inches(0.8), Inches(5.75), Inches(3.9), Inches(0.9), share.upper(), size=16, color=EMBER, bold=True, font=TITLE_FONT,
             spacing=150, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(5.2), Inches(1.7), Inches(7.7), Inches(0.9), [f"**WHO**   {who}", f"**GOAL**   {goal}"], size=15)
        text(s, Inches(5.2), Inches(2.75), Inches(3), Inches(0.35), "ACTIONS", size=12, color=RED, bold=True, font=TITLE_FONT, spacing=250)
        bullets(s, Inches(5.2), Inches(3.15), Inches(7.7), Inches(2.3), actions, size=15, gap=6)
        card(s, Inches(5.2), Inches(5.35), Inches(3.75), Inches(1.45), "KPIs", kpis, accent=EMBER, size=12)
        card(s, Inches(9.15), Inches(5.35), Inches(3.75), Inches(1.45), "Guardrail", guardrail, accent=RED, size=12)
        notes(s, f"{name}: {who.lower()}. Goal: {goal.lower()}. Guardrail: {guardrail.lower()}.")

    # 19. Responsible AI
    s = new()
    chrome(s, 19, "Responsible AI and fairness", "THE SAFEGUARDS")
    bullets(s, Inches(0.55), Inches(1.7), Inches(6.2), Inches(1.5), [
        "**Behaviour-only clustering:** demographics excluded when forming segments",
        "**Proxy check (Cramér's V)** against demographics that were *not* used",
    ], size=15)
    table(s, Inches(0.55), Inches(3.0), Inches(6.2), [
        ["Attribute", "Cramér's V", "Review needed"],
        ["Parental status", "0.53", "YES"],
        ["Age band", "0.14", "No"],
        ["Education", "0.14", "No"],
        ["Marital status", "0.06", "No"],
    ], col_widths=[2, 1, 1.2], size=13, row_h=0.45, bold_rows=(1,))
    bullets(s, Inches(0.55), Inches(5.4), Inches(6.2), Inches(1.5), [
        "Never use segments for **credit, pricing or eligibility** decisions",
        "Guardrails: contact caps, no dark patterns, affordability limits, transparent names",
    ], size=14, gap=5)
    picture(s, FIG / "13_demographics.png", Inches(7.1), Inches(1.7), Inches(5.8), Inches(5.1))
    notes(s, "Families naturally have tighter budgets, so the behaviour is genuine. The risk is in how segments are used; "
             "this check re-runs on every retrain.")

    # 20. Website
    s = new()
    chrome(s, 20, "Production website", "THE COMMAND CENTRE")
    bullets(s, Inches(0.55), Inches(1.7), Inches(4.6), Inches(5.2), [
        "**Private workspace** per account; cloud PostgreSQL logins",
        "**Overview, Profiles, Explorer, Assign, Methodology** for any dataset",
        "**Data agent (Gemini):** upload CSV/Excel/JSON up to 1 lakh rows; it cleans, trains and rebuilds every page",
        "**Excel reports with charts**, chat history, REST API at /api/v1",
        "**My account:** name, password, workspace, recent chats",
        "**Premium boutique design**, light/dark, Calm mode",
    ], size=13.5, gap=7)
    shots = ["site_overview.png", "site_segments.png", "site_agent.png", "site_account.png"]
    for i, shot in enumerate(shots):
        col, row = i % 2, i // 2
        picture(s, FIG / shot, Inches(5.45 + col * 3.75), Inches(1.75 + row * 2.6), Inches(3.55), Inches(2.4))
    text(s, Inches(5.45), Inches(6.72), Inches(7.4), Inches(0.3), "LIVE: customeriq.up.railway.app", size=10, color=EMBER, bold=True,
         font=TITLE_FONT, spacing=200)
    notes(s, "Upload any table and every page switches to that data for your account. The customer model, the dataset workspaces "
             "and the notebook share one package, so what users see always matches the analysis.")

    # 21. Architecture
    s = new()
    chrome(s, 21, "Solution architecture", "THE MACHINE")
    top = ["Raw / messy\nfile", "Load &\nparse", "Clean\n(audit)", "Features", "Preprocess", "Select k +\nK-Means", "Profiles &\nplaybooks"]
    bw, gap_x, y1 = Inches(1.5), Inches(0.31), Inches(1.85)
    for i, label in enumerate(top):
        x = Inches(0.55) + i * (bw + gap_x)
        rect(s, x, y1, bw, Inches(0.95), fill=RED if i == 0 else PANEL, line=RED)
        text(s, x, y1, bw, Inches(0.95), label, size=12, color=WHITE, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(top) - 1:
            arrow(s, x + bw, y1 + Inches(0.475), x + bw + gap_x, y1 + Inches(0.475))
    store_x = Inches(0.55) + 6 * (bw + gap_x)
    y2 = Inches(3.55)
    bottom = [("Artifacts &\nuser workspaces", PANEL), ("FastAPI\n/api/v1", RED), ("Website\n(SPA)", PANEL), ("Data agent\n+ Gemini", PANEL),
              ("PostgreSQL\naccounts", PANEL), ("CRM / other\nsystems", PANEL)]
    bw2, gap2 = Inches(1.75), Inches(0.37)
    for i, (label, color) in enumerate(bottom):
        x = Inches(0.55) + i * (bw2 + gap2)
        rect(s, x, y2, bw2, Inches(0.95), fill=color, line=RED)
        text(s, x, y2, bw2, Inches(0.95), label, size=12, color=WHITE, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, store_x + bw / 2, y1 + Inches(0.95), Inches(0.55) + bw2 / 2, y2)
    arrow(s, Inches(0.55) + bw2, y2 + Inches(0.475), Inches(0.55) + bw2 + gap2, y2 + Inches(0.475))
    # The API serves the website, the agent, the account database and outside systems through one "bus".
    center = lambda i: Inches(0.55) + i * (bw2 + gap2) + bw2 // 2  # noqa: E731
    bus_y = y2 + Inches(1.2)
    arrow(s, center(1), y2 + Inches(0.95), center(1), bus_y, head=False)
    arrow(s, center(1), bus_y, center(5), bus_y, head=False)
    for i in range(2, 6):
        arrow(s, center(i), bus_y, center(i), y2 + Inches(0.95))
    bullets(s, Inches(0.55), Inches(5.0), Inches(12.35), Inches(1.9), [
        "**One package** (segmentation/) drives the notebook, CLI, API, website and dashboard",
        "**Hardened:** CSP and security headers, rate limits, upload limits, formula-injection protection, scrypt passwords, HttpOnly sessions, non-root Docker",
        "**72 automated tests:** unit, API, security, accounts, messy-data cleaning, dataset workspaces, AI tool calls",
    ], size=14, gap=6)
    notes(s, "One reusable package drives every surface, so training and scoring always use identical transformations.")

    # 22. Conclusions
    s = new()
    chrome(s, 22, "Conclusions and next steps", "THE AFTERMATH")
    card(s, Inches(0.55), Inches(1.75), Inches(6.0), Inches(5.0), "What we proved", [
        "**Four stable, behaviour-based segments**, k justified by elbow, silhouette, Davies-Bouldin and stability",
        "**Two premium segments = 78% of revenue**, but they need **opposite** strategies",
        "**Digital Deal-Seekers** are the main growth opportunity",
        "**Budget-Conscious Browsers** need low-cost everyday value",
        "The agent **survives a messy 1 lakh export** and still finds the same segments",
    ], size=14)
    card(s, Inches(6.9), Inches(1.75), Inches(6.0), Inches(5.0), "Next steps", [
        "A/B-test each playbook against a control group",
        "Add transaction-level data for true frequency and order-level recency",
        "Monitor drift monthly (shares, centroid movement); retrain quarterly",
        "Re-run the fairness check on every retrain",
    ], accent=EMBER, size=14)
    notes(s, "Four stable segments, clear playbooks and guardrails, and a production system that keeps them up to date.")

    # 23. Closing
    s = new()
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0, 0, Inches(0.12), H, fill=RED)
    rect(s, 0, Inches(7.38), W, Inches(0.12), fill=RED)
    doomsday_clock(s, Inches(3.4), Inches(3.75), Inches(2.0))
    text(s, Inches(6.6), Inches(1.3), Inches(6.3), Inches(1.4), "00:00", size=96, color=RED, bold=True, font=TITLE_FONT, spacing=300)
    text(s, Inches(6.6), Inches(2.85), Inches(6.3), Inches(0.8), "MIDNIGHT FOR GUESSWORK.", size=30, color=WHITE, bold=True, font=TITLE_FONT,
         spacing=150)
    text(s, Inches(6.6), Inches(3.65), Inches(6.3), Inches(0.9), "Every customer now has a segment, a playbook and a guardrail.", size=18,
         color=ASH)
    text(s, Inches(6.6), Inches(4.75), Inches(6.3), Inches(0.6), "QUESTIONS?", size=26, color=EMBER, bold=True, font=TITLE_FONT, spacing=400)
    text(s, Inches(6.6), Inches(5.6), Inches(6.3), Inches(0.9), [
        "Live site  ·  customeriq.up.railway.app",
        "Code  ·  github.com/uditmehrotradkd-cpu/customer-iq",
    ], size=13, color=DIM, font=TITLE_FONT, spacing=100)
    text(s, Inches(6.6), Inches(6.6), Inches(6.3), Inches(0.4), "TEAM DOOMSDAY", size=14, color=RED, bold=True, font=TITLE_FONT, spacing=600)
    notes(s, "Thank you. Questions?")

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads" / "Doomsday_CustomerIQ.pptx"
    print(f"Saved {build(target)}")

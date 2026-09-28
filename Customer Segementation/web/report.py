"""Excel reports with native charts, so a downloaded result explains itself without any analysis tools."""
from __future__ import annotations

import io
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from segmentation.config import CHANNEL_PURCHASE_COLUMNS, SPEND_COLUMNS

from .services import CSV_FORMULA_PREFIXES

# Same palette as the website's segment colours.
PALETTE = ["2A9D8F", "E76F51", "6C63FF", "E9C46A", "8AB17D", "F4A261", "264653", "B56576", "457B9D", "9C6644"]
HEADER_FILL = PatternFill("solid", fgColor="312E81")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=18, color="312E81")
MAX_METRICS = 6
SKIP_COLUMNS = {"Segment", "Distance_To_Centroid", "Assignment_Margin", "ID"}
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WORKSPACE_METRICS = ["Income", "Total_Spend", "Total_Purchases", "Avg_Order_Value", "Recency", "Deal_Ratio"]
SCORED_METRICS = ["Income", "Total_Spend", "Total_Purchases", "Recency", "NumWebVisitsMonth", "NumDealsPurchases"]


def _safe(value):
    """Uploaded text must never become an Excel formula."""
    if isinstance(value, np.generic):
        value = value.item()
    return f"'{value}" if isinstance(value, str) and value.startswith(CSV_FORMULA_PREFIXES) else value


def _label(name: str) -> str:
    return _safe(str(name).replace("_", " "))


def _pick_metrics(data: pd.DataFrame, segment_col: str) -> list[str]:
    """Numeric columns whose averages differ most between segments (largest between-segment spread)."""
    numeric = [c for c in data.select_dtypes(include="number").columns if c not in SKIP_COLUMNS and data[c].nunique() > 2]
    scores = {}
    for col in numeric:
        std = data[col].std()
        if std and np.isfinite(std):
            means = data.groupby(segment_col)[col].mean()
            scores[col] = float((means.max() - means.min()) / std)
    return sorted(scores, key=scores.get, reverse=True)[:MAX_METRICS]


def _header(ws, row: int, values: list) -> None:
    for i, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _colour_points(series, n: int) -> None:
    for i in range(n):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
        point.graphicalProperties.line.solidFill = PALETTE[i % len(PALETTE)]
        series.dPt.append(point)


def _labels(**show) -> DataLabelList:
    labels = DataLabelList()
    labels.showVal = labels.showPercent = labels.showSerName = labels.showCatName = labels.showLegendKey = labels.showLeaderLines = False
    for key, value in show.items():
        setattr(labels, key, value)
    return labels


def _bar(ws, title: str, data_ref: Reference, cats: Reference, n: int, number_format: str = "#,##0", zero_base: bool = True) -> BarChart:
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.y_axis.majorGridlines = None
    chart.y_axis.numFmt = number_format
    # Bars must start at zero, otherwise small differences look dramatic.
    if zero_base:
        chart.y_axis.scaling.min = 0
    chart.legend = None
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 7.5, 15
    series = chart.series[0]
    _colour_points(series, n)
    series.dLbls = _labels(showVal=True)
    series.dLbls.numFmt = number_format
    series.dLbls.position = "outEnd"
    # openpyxl hides axes by default in recent versions; show them explicitly.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    return chart


def build_report(
    data: pd.DataFrame,
    title: str,
    source: str,
    segment_col: str = "Segment_Name",
    metrics: list[str] | None = None,
    notes: dict[str, list[str]] | None = None,
) -> bytes:
    """Return an .xlsx workbook: summary charts, per-segment comparisons, plain-language insights and the data."""
    metrics = [m for m in (metrics or _pick_metrics(data, segment_col)) if m in data.columns][:MAX_METRICS]
    counts = data[segment_col].value_counts()
    order = list(counts.index)
    n_seg = len(order)
    means = data.groupby(segment_col)[metrics].mean().reindex(order) if metrics else pd.DataFrame(index=order)

    wb = Workbook()

    # ---------------------------------------------------------------- summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = _safe(title)
    ws["A1"].font = TITLE_FONT
    ws["A2"] = _safe(f"Source: {source}  \u00b7  {len(data):,} rows  \u00b7  {n_seg} segments  \u00b7  generated {datetime.now():%d %b %Y %H:%M}")
    ws["A2"].font = Font(italic=True, color="6B7280")
    _header(ws, 4, ["Segment", "Customers", "Share of customers"])
    for i, name in enumerate(order, start=5):
        ws.cell(row=i, column=1, value=_safe(name)).font = Font(bold=True, color=PALETTE[(i - 5) % len(PALETTE)])
        ws.cell(row=i, column=2, value=int(counts[name])).number_format = "#,##0"
        ws.cell(row=i, column=3, value=float(counts[name] / len(data))).number_format = "0.0%"
    last = 4 + n_seg
    ws.cell(row=last + 1, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=last + 1, column=2, value=len(data)).number_format = "#,##0"
    ws.cell(row=last + 1, column=3, value=1).number_format = "0.0%"

    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    ws.add_chart(_bar(ws, "How many customers are in each segment", Reference(ws, min_col=2, min_row=4, max_row=last), cats, n_seg), "E4")
    pie = PieChart()
    pie.title = "Share of customers"
    pie.add_data(Reference(ws, min_col=3, min_row=4, max_row=last), titles_from_data=True)
    pie.set_categories(cats)
    pie.height, pie.width = 9, 17
    pie.legend.position = "r"
    _colour_points(pie.series[0], n_seg)
    pie.series[0].dLbls = _labels(showPercent=True)
    pie.series[0].dLbls.position = "bestFit"
    ws.add_chart(pie, "E20")

    guide_row = last + 4
    ws.cell(row=guide_row, column=1, value="How to read this report").font = Font(bold=True, size=13)
    guide = [
        "1. This sheet: how big each segment is.",
        "2. 'Segment profiles': the average of the most telling measures for each segment, with a chart per measure.",
        "3. 'Segment insights': what makes each segment different, in plain words.",
        "4. 'Data': every row of your file with the segment it belongs to (use the filter arrows).",
    ]
    for j, text in enumerate(guide, start=1):
        ws.cell(row=guide_row + j, column=1, value=text)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18

    # ---------------------------------------------------------------- profiles
    if metrics:
        wp = wb.create_sheet("Segment profiles")
        wp["A1"] = "Segment profiles: averages per segment"
        wp["A1"].font = Font(bold=True, size=14, color="312E81")
        wp["A2"] = "Compare each segment with the 'All rows' line to see what is above or below normal."
        wp["A2"].font = Font(italic=True, color="6B7280")
        _header(wp, 4, ["Segment", "Customers", *[f"Avg {_label(m)}" for m in metrics]])
        formats = {m: "#,##0" if abs(float(data[m].mean())) >= 100 else "0.00" for m in metrics}
        for i, name in enumerate(order, start=5):
            wp.cell(row=i, column=1, value=_safe(name)).font = Font(bold=True)
            wp.cell(row=i, column=2, value=int(counts[name])).number_format = "#,##0"
            for j, m in enumerate(metrics, start=3):
                value = means.loc[name, m]
                wp.cell(row=i, column=j, value=None if pd.isna(value) else round(float(value), 2)).number_format = formats[m]
        overall = 5 + n_seg
        wp.cell(row=overall, column=1, value="All rows").font = Font(italic=True, bold=True)
        wp.cell(row=overall, column=2, value=len(data)).number_format = "#,##0"
        for j, m in enumerate(metrics, start=3):
            wp.cell(row=overall, column=j, value=round(float(data[m].mean()), 2)).number_format = formats[m]
        wp.column_dimensions["A"].width = 42
        for j in range(2, len(metrics) + 3):
            wp.column_dimensions[get_column_letter(j)].width = 16
        seg_cats = Reference(wp, min_col=1, min_row=5, max_row=4 + n_seg)
        for k, m in enumerate(metrics):
            chart = _bar(wp, f"Average {_label(m)} by segment", Reference(wp, min_col=3 + k, min_row=4, max_row=4 + n_seg), seg_cats, n_seg, formats[m], zero_base=bool(means[m].min() >= 0))
            wp.add_chart(chart, f"{'A' if k % 2 == 0 else 'J'}{overall + 3 + (k // 2) * 16}")

    # ---------------------------------------------------------------- insights
    if notes:
        wi = wb.create_sheet("Segment insights")
        wi["A1"] = "What makes each segment different"
        wi["A1"].font = Font(bold=True, size=14, color="312E81")
        row = 3
        for i, name in enumerate(order):
            wi.cell(row=row, column=1, value=_safe(f"{name}  ({counts[name] / len(data):.0%} of rows)")).font = Font(bold=True, size=12, color=PALETTE[i % len(PALETTE)])
            row += 1
            for line in notes.get(name, []) or ["Close to the overall average on every measure."]:
                cell = wi.cell(row=row, column=1, value=f"\u2022 {line}")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                row += 1
            row += 1
        wi.column_dimensions["A"].width = 110

    # ---------------------------------------------------------------- data
    wd = wb.create_sheet("Data")
    columns = [segment_col, *[c for c in data.columns if c != segment_col]]
    _header(wd, 1, [_label(c) for c in columns])
    values = data[columns].astype(object).where(data[columns].notna(), None)
    for record in values.itertuples(index=False, name=None):
        wd.append([_safe(v) for v in record])
    wd.freeze_panes = "B2"
    wd.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(data) + 1}"
    wd.column_dimensions["A"].width = 36
    for j in range(2, len(columns) + 1):
        wd.column_dimensions[get_column_letter(j)].width = 14

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def recommendation_notes(service) -> dict[str, list[str]]:
    notes = {}
    for r in service.segmenter.recommendations_:
        lines = [*r.get("traits", []), r["summary"] if not r.get("traits") else None, f"Goal: {r['goal']}"]
        if r.get("preferred_channel"):
            lines.append(f"Best channel: {r['preferred_channel']}")
        notes[r["name"]] = [line for line in lines if line] + [f"Next step: {a}" for a in r["actions"]]
    return notes


def scored_report(service, scored: pd.DataFrame, name: str) -> bytes:
    """Report for a file whose rows were assigned to the workspace's segments."""
    data = scored.drop(columns="Segment", errors="ignore")
    if service.mode == "generic":
        metrics = [m for m in service.metrics if m in data.columns and pd.api.types.is_numeric_dtype(data[m])] or None
        return build_report(data, f"Segment assignment for {name}", name, metrics=metrics, notes=recommendation_notes(service))
    data = data.assign(
        Total_Spend=data[[c for c in SPEND_COLUMNS if c in data]].apply(pd.to_numeric, errors="coerce").sum(axis=1),
        Total_Purchases=data[[c for c in CHANNEL_PURCHASE_COLUMNS if c in data]].apply(pd.to_numeric, errors="coerce").sum(axis=1),
    )
    return build_report(data, f"Segment assignment for {name}", name, metrics=SCORED_METRICS, notes=recommendation_notes(service))


def workspace_report(service, source: str) -> bytes:
    """Report for the model behind every page: its rows, segment averages and next steps."""
    rows = service.customers.drop(columns=["Segment"], errors="ignore")
    metrics = service.metrics if service.mode == "generic" else WORKSPACE_METRICS
    title = f"Segments in {source}" if service.mode == "generic" else "Customer segmentation report"
    return build_report(rows, title, source, metrics=metrics, notes=recommendation_notes(service))

"""Agent Q&A knowledge base and multi-format (CSV / Excel / JSON) file handling."""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from segmentation import SegmentationConfig  # noqa: E402
from segmentation.config import ARTIFACTS_DIR, MODEL_FILENAME  # noqa: E402
from web.agent import handle_message  # noqa: E402
from web.file_io import FileFormatError, analyze_table, normalize_columns, read_table  # noqa: E402
from web.main import create_app  # noqa: E402
from web.registry import ModelRegistry  # noqa: E402
from web.report import build_report  # noqa: E402
from web.services import SegmentationService  # noqa: E402
from web.settings import Settings  # noqa: E402

pytestmark = pytest.mark.skipif(not (ARTIFACTS_DIR / MODEL_FILENAME).exists(), reason="Run `python run_pipeline.py train` first")
FAST = SegmentationConfig(k_max=5, n_init=3, stability_runs=2)


@pytest.fixture(scope="module")
def service() -> SegmentationService:
    return SegmentationService.load(ARTIFACTS_DIR, auto_train=False)


@pytest.fixture()
def registry(service, tmp_path) -> ModelRegistry:
    return ModelRegistry(service, tmp_path, FAST)


@pytest.fixture(scope="module")
def raw(service) -> pd.DataFrame:
    return service.customers[[c for c in service.customers.columns if c not in ("Segment", "Segment_Name")]].head(300)


def ask(text, registry):
    return handle_message(text, registry, None)


def test_questions_are_answered(registry, service):
    k = service.segmenter.n_segments
    assert f"{k} segments" in ask("Why 4 segments?", registry)["reply"]
    assert "silhouette" in ask("What is the silhouette score?", registry)["reply"].lower()
    assert "Digital Deal-Seekers" in ask("Tell me about digital deal-seekers", registry)["reply"]
    assert "highest median income" in ask("Which segment has the highest income?", registry)["reply"]
    assert "Comparison" in ask("compare premium and affluent", registry)["reply"]
    assert "Excel" in ask("What file formats can I upload?", registry)["reply"]
    assert "discount" in ask("what does deal ratio mean?", registry)["reply"]
    assert "Explorer" in ask("what is on the explorer page?", registry)["reply"]
    assert "Cram" in ask("is the model fair?", registry)["reply"]
    assert "after cleaning" in ask("how was the data cleaned?", registry)["reply"]
    for text in ("how do I train on my own data?", "how do I publish a model?"):
        res = ask(text, registry)
        assert res["intent"] == "answer", text
    assert ask("what is the weather today?", registry)["intent"] == "unknown"


def test_answers_include_links(registry):
    res = ask("How was the number of segments chosen?", registry)
    assert res["intent"] == "answer"
    assert any(href.startswith("#/") for _, href in res["links"])


def test_commands_still_route_to_actions(registry):
    assert ask("can you give me premium customers with income over 70k?", registry)["intent"] == "filtered"
    assert ask("please generate 200 synthetic customers", registry)["intent"] == "synthetic"
    assert ask("train with 3 segments", registry)["intent"] == "train"


@pytest.mark.parametrize("fmt", ["csv", "tsv", "xlsx", "json", "json_wrapped", "jsonl"])
def test_read_table_formats(raw, fmt):
    frame = raw.head(20)
    if fmt == "csv":
        content, name = frame.to_csv(index=False).encode(), "a.csv"
    elif fmt == "tsv":
        content, name = frame.to_csv(index=False, sep="\t").encode(), "a.tsv"
    elif fmt == "xlsx":
        buffer = io.BytesIO()
        frame.to_excel(buffer, index=False)
        content, name = buffer.getvalue(), "a.xlsx"
    elif fmt == "json":
        content, name = frame.to_json(orient="records", date_format="iso").encode(), "a.json"
    elif fmt == "json_wrapped":
        content, name = json.dumps({"customers": json.loads(frame.to_json(orient="records", date_format="iso"))}).encode(), "a.json"
    else:
        content, name = frame.to_json(orient="records", lines=True, date_format="iso").encode(), "a.jsonl"
    parsed = read_table(content, name, 1000)
    assert len(parsed) == 20
    assert {"ID", "Income", "MntWines"} <= set(parsed.columns)


def test_read_table_rejects_bad_files():
    with pytest.raises(FileFormatError):
        read_table(b"%PDF-1.4", "x.pdf", 100)
    with pytest.raises(FileFormatError):
        read_table(b"{not json", "x.json", 100)
    with pytest.raises(FileFormatError):
        read_table(b"a,b\n" + b"1,2\n" * 50, "x.csv", 10)


def test_normalize_and_analyze(raw):
    frame = raw.head(10).rename(columns={"Year_Birth": "birth year", "Income": "income", "ID": "customer_id"})
    normalized, renames = normalize_columns(frame)
    assert {"Year_Birth", "Income", "ID"} <= set(normalized.columns)
    report = analyze_table(normalized, renames)
    assert report["compatible"] and report["rows"] == 10
    broken = analyze_table(normalized.drop(columns=["Income"]), {})
    assert not broken["compatible"] and "Income" in broken["missing_columns"]


def _sign_in(client) -> None:
    res = client.post("/api/v1/auth/register", json={"username": "tester", "name": "Tester", "password": "s3cret-pass"})
    assert res.status_code == 200, res.text


def _payload(dataset: dict) -> dict:
    return {k: v for k, v in dataset.items() if k != "filename"}


def test_report_never_writes_formulas():
    frame = pd.DataFrame({"Segment_Name": ["=cmd|'/c calc'!A1", "B"] * 10, "Note": ["=HYPERLINK(\"x\")", "ok"] * 10, "Value": range(20)})
    workbook = load_workbook(io.BytesIO(build_report(frame, "t", "s")))
    cells = [c for ws in workbook.worksheets for row in ws.iter_rows() for c in row]
    assert not any(c.data_type == "f" for c in cells)
    assert any(str(c.value).startswith("'=HYPERLINK") for c in cells)


def test_agent_file_endpoint(service, raw, tmp_path):
    settings = Settings(allowed_hosts=["testserver"], rate_limit_per_minute=1000, model_store_dir=tmp_path, min_train_rows=100)
    app = create_app(settings, service=service)
    buffer = io.BytesIO()
    raw.to_excel(buffer, index=False)
    xlsx = ("mine.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with TestClient(app) as client:
        app.state.registry.config = FAST
        _sign_in(client)
        analyzed = client.post("/api/v1/agent/file", files={"file": xlsx}, data={"action": "analyze"})
        assert analyzed.status_code == 200, analyzed.text
        assert analyzed.json()["file_actions"] == ["auto", "score", "cluster"]

        scored = client.post("/api/v1/agent/file", files={"file": xlsx}, data={"action": "score"})
        assert scored.status_code == 200, scored.text
        report = client.post("/api/v1/datasets", json=_payload(scored.json()["dataset"]))
        assert report.headers["content-type"].startswith("application/vnd.openxmlformats")
        workbook = load_workbook(io.BytesIO(report.content))
        assert workbook.sheetnames == ["Summary", "Segment profiles", "Segment insights", "Data"]
        assert len(workbook["Summary"]._charts) == 2 and len(workbook["Segment profiles"]._charts) >= 3
        assert workbook["Data"].max_row == len(raw) + 1
        dataset = scored.json()["downloads"][0]["dataset"]
        assert dataset["kind"] == "result"
        result = pd.read_csv(io.BytesIO(client.post("/api/v1/datasets", json=_payload(dataset)).content))
        assert len(result) == len(raw) and "Segment_Name" in result.columns
        assert client.post("/api/v1/datasets", json={"kind": "report"}).status_code == 200

        jsonl = ("mine.jsonl", raw.to_json(orient="records", lines=True, date_format="iso").encode(), "application/json")
        trained = client.post("/api/v1/agent/file", files={"file": jsonl}, data={"action": "train", "k": "3"})
        assert trained.status_code == 200, trained.text
        assert trained.json()["draft"]["k"] == 3

        assert client.post("/api/v1/datasets", json={"kind": "result", "token": "../../etc"}).status_code == 422
        assert client.post("/api/v1/datasets", json={"kind": "result", "token": "a" * 32}).status_code == 422
        bad = client.post("/api/v1/agent/file", files={"file": ("x.pdf", b"%PDF", "application/pdf")}, data={"action": "analyze"})
        assert bad.status_code == 422
        assert client.post("/api/v1/agent/file", files={"file": xlsx}, data={"action": "delete"}).status_code == 422


def test_file_without_enrolment_date_scores_and_trains(service, raw, tmp_path):
    settings = Settings(allowed_hosts=["testserver"], rate_limit_per_minute=1000, model_store_dir=tmp_path, min_train_rows=100)
    app = create_app(settings, service=service)
    no_date = ("kyc.csv", raw.drop(columns=["ID", "Dt_Customer"]).to_csv(index=False).encode(), "text/csv")
    with TestClient(app) as client:
        app.state.registry.config = FAST
        _sign_in(client)
        analyzed = client.post("/api/v1/agent/file", files={"file": no_date}, data={"action": "analyze"})
        assert analyzed.status_code == 200 and "score" in analyzed.json()["file_actions"]
        scored = client.post("/api/v1/agent/file", files={"file": no_date}, data={"action": "score"})
        assert scored.status_code == 200, scored.text
        result = pd.read_csv(io.BytesIO(client.post("/api/v1/datasets", json=_payload(scored.json()["downloads"][0]["dataset"])).content))
        assert len(result) == len(raw) and "Dt_Customer" not in result.columns
        trained = client.post("/api/v1/agent/file", files={"file": no_date}, data={"action": "train", "k": "3"})
        assert trained.status_code == 200, trained.text


def test_cluster_any_dataset(service, tmp_path):
    settings = Settings(allowed_hosts=["testserver"], rate_limit_per_minute=1000, model_store_dir=tmp_path, min_train_rows=100)
    app = create_app(settings, service=service)
    rng = np.random.default_rng(0)
    n = 300
    group = rng.integers(0, 3, n)
    frame = pd.DataFrame(
        {
            "customer_id": np.arange(n),
            "Full_Name": [f"Person {i}" for i in range(n)],
            "Balance": np.array([1_000, 20_000, 90_000])[group] * rng.uniform(0.8, 1.2, n),
            "Transactions": np.array([5, 40, 12])[group] + rng.integers(0, 3, n),
            "Region": np.array(["North", "South", "East"])[group],
            "Opened": pd.date_range("2020-01-01", periods=n, freq="D").strftime("%Y-%m-%d"),
        }
    )
    upload = ("accounts.csv", frame.to_csv(index=False).encode(), "text/csv")
    with TestClient(app) as client:
        _sign_in(client)
        analyzed = client.post("/api/v1/agent/file", files={"file": upload}, data={"action": "analyze"})
        assert analyzed.json()["file_actions"] == ["auto"]
        for action in ("cluster", "train"):
            res = client.post("/api/v1/agent/file", files={"file": upload}, data={"action": action})
            assert res.status_code == 200, res.text
            body = res.json()
            assert body["intent"] == "clustered" and len(body["segments_view"]) >= 2
        assert "Full_Name (personal data" in body["reply"] and "(identifier)" in body["reply"]
        assert body["dataset"]["kind"] == "report"
        labelled = pd.read_csv(io.BytesIO(client.post("/api/v1/datasets", json=_payload(body["downloads"][0]["dataset"])).content))
        assert len(labelled) == n and {"Segment", "Segment_Name"} <= set(labelled.columns)
        assert pd.crosstab(labelled["Segment"], frame["Region"]).gt(0).sum(axis=1).eq(1).all()

        # The whole site now describes this dataset.
        assert body["applied"] and client.get("/api/v1/health").json()["mode"] == "generic"
        overview = client.get("/api/v1/overview").json()
        assert overview["labels"]["dataset"] == "accounts.csv" and overview["kpis"]["customers"] == n
        assert overview["labels"]["value"] == "Balance"
        for path in ("/pca", "/pca3d", "/segments", "/segments/0", "/fingerprint", "/diagnostics", "/customers/export"):
            assert client.get(f"/api/v1{path}").status_code == 200, path
        features = client.get("/api/v1/explorer/features").json()
        assert "Balance" in features["numeric"] and "Region" in features["categorical"]
        assert "Full_Name" not in features["numeric"] + features["categorical"]
        assert client.get("/api/v1/explorer/demographics", params={"attribute": "Region"}).status_code == 200
        form = client.get("/api/v1/assign/defaults").json()
        assert form["mode"] == "generic" and {f["name"] for f in form["fields"]} >= {"Balance", "Region", "Opened"}
        assigned = client.post("/api/v1/assign/generic", json={"values": {"Balance": 90_000, "Transactions": 12, "Region": "East", "Opened": "2020-03-01"}})
        assert assigned.status_code == 200 and assigned.json()["segment_name"] in {s["name"] for s in overview["segments"]}
        assert client.post("/api/v1/assign", json={}).status_code in (409, 422)
        chat = client.post("/api/v1/agent/message", json={"message": "which segment has the highest balance?"}).json()
        assert chat["intent"] == "answer" and "Balance" in chat["reply"]
        assert client.post("/api/v1/agent/message", json={"message": "generate 100 synthetic customers"}).status_code in (200, 422)

        scored = client.post("/api/v1/agent/file", files={"file": upload}, data={"action": "score"})
        assert scored.status_code == 200, scored.text
        rescored = pd.read_csv(io.BytesIO(client.post("/api/v1/datasets", json=_payload(scored.json()["downloads"][0]["dataset"])).content))
        assert (rescored["Segment"] == labelled["Segment"]).all()

        assert client.post("/api/v1/agent/message", json={"message": "reset to original"}).json()["intent"] == "reset_done"
        assert client.get("/api/v1/health").json()["mode"] == "customer"
